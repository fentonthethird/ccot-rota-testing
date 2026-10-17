"""End to end through the admin workflow on a temporary data folder."""
import datetime as dt
import json
from urllib.parse import parse_qs, urlparse

import openpyxl

from ccot import allocate, codec, costing, dashboard, importer, xlsx_export, xlsx_import
from ccot import calendar as cal
from ccot import roster as R
from ccot import windows as W
from ccot import years as Y
from conftest import D, add_people

PEOPLE = {"AW": ("Alan Williams", 10, 2), "XW": ("Ximena Watson", 10, 2), "AJ": ("Ankit Jain", 12, 2), "DL": ("Dagan Lonsdale", 5, 0), "LDC": ("Leonel De Caires", 0, 0)}


def make_window(store):
    add_people(store, PEOPLE)
    return W.create(store, "2026-27-Q4", D("2027-01-01"), D("2027-03-31"), D("2026-11-30"))


def code_for(store, w, ini, picker, pairs=()):
    start = D(w["start"])
    n = W.length(w)
    days = [1 if picker(start + dt.timedelta(days=i)) else 0 for i in range(n)]
    return codec.encode(ini, start, days, pairs, token=w["tokens"][ini])


def test_links_carry_identity_window_and_shifts_owed(store):
    w = make_window(store)
    y = store.year("2026-27")
    Y.set_shift(y, D("2026-05-05"), "AW", Y.JP, "import")
    Y.set_shift(y, D("2026-05-09"), "AW", Y.JP, "import")
    store.save_year(y)
    links = {l["initials"]: l["url"] for l in W.links(store, w["id"])}
    q = {k: v[0] for k, v in parse_qs(urlparse(links["AW"]).query).items()}
    assert q == {"c": "AW", "k": w["tokens"]["AW"], "n": "Alan", "w": "2026-27-Q4", "s": "2027-01-01", "e": "2027-03-31", "d": "2026-11-30", "jw": "9", "je": "1"}
    assert links["AW"].startswith("https://example.github.io/ccot-rota/?")
    assert W.windows_left(w) == 1 and W.default_allowance(store, w, R.get(store, "AW")) == {"weekday": 9, "weekend": 1}
    q1 = W.create(store, "2027-28-Q1", D("2027-04-01"), D("2027-06-30"))
    assert W.windows_left(q1) == 4


def test_import_status_heatmap_allocate_accept_and_publish(store, tmp_path):
    w = make_window(store)
    aw = code_for(store, w, "AW", lambda d: d.weekday() < 5 and d.day % 2 == 0)
    xw = code_for(store, w, "XW", lambda d: True, pairs=[s for s in codec.pair_slots(D(w["start"]), W.length(w))][:3])
    aj = code_for(store, w, "AJ", lambda d: d.weekday() >= 5 or d.day % 3 == 0)
    dl_old = code_for(store, w, "DL", lambda d: d.month == 1)
    dl_new = code_for(store, w, "DL", lambda d: d.month == 2)
    junk = aw[:-4] + "ZZZZ"
    text = "AW says: %s\n\nXimena: %s\nAnkit %s\n%s\n" % (aw, xw, aj, dl_old)
    rep = importer.import_text(store, w["id"], text)
    assert [r["initials"] for r in rep["accepted"]] == ["AW", "XW", "AJ", "DL"] and not rep["rejected"]
    rep = importer.import_text(store, w["id"], dl_new + "\n" + junk)
    assert [r["initials"] for r in rep["replaced"]] == ["DL"] and len(rep["rejected"]) == 1
    assert any(word in rep["rejected"][0]["error"] for word in ("corrupt", "checksum", "long", "truncated"))
    # a code made from a different window is refused
    other = codec.encode("AW", D("2026-10-01"), [1] * 92, token=w["tokens"]["AW"])
    rep = importer.import_text(store, w["id"], other)
    assert "different window" in rep["rejected"][0]["error"]
    st = importer.status(store, w["id"])
    assert [x["initials"] for x in st["submitted"]] == ["AW", "XW", "AJ", "DL"] and [x["initials"] for x in st["missing"]] == ["LDC"]
    rows = importer.heatmap(store, w["id"])
    assert len(rows) == 90 and rows[0]["type"] == "bh" and "XW" in rows[0]["offers"]
    assert "0 dates with no offers" in importer.heatmap_text(rows)  # XW offered every date, so nothing is uncovered
    html = importer.heatmap_html(rows, "t")
    assert "01 Jan" in html and "BH" in html

    draft = allocate.run(store, w["id"], time_limit=30)
    assert draft["unfillable"] == [] and all(s["optimal"] for s in draft["stages"])
    assert draft["fill"]["weekday"]["filled"] == draft["fill"]["weekday"]["total"]
    c = draft["consultants"]
    assert c["AW"]["jp_window"]["weekday"] == 10 and c["AJ"]["jp_window"] == {"weekday": 12, "weekend": 2}
    assert c["DL"]["jp_window"]["weekday"] == 5 and c["LDC"]["submitted"] is False
    assert c["XW"]["locum_window"] > 0
    # the first shifts of each type are job plan, the rest locum, and 1 Jan is a bank holiday shift
    assert draft["kinds"]["2027-01-01"] == "BH"
    xw_dates = sorted(k for k, v in draft["assignment"].items() if v == "XW" and cal.day_type(D(k)) == cal.WEEKDAY)
    kinds = [draft["kinds"][k] for k in xw_dates]
    assert kinds[:10] == ["JP"] * 10 and set(kinds[10:]) <= {"LOCUM"}

    allocate.accept(store, w["id"])
    y = store.year("2026-27")
    assert len(y["shifts"]) == 90 and y["unfilled"] == [] and y["allocated_windows"] == ["2026-27-Q4"]
    assert store.window(w["id"])["status"] == "allocated"
    # re-running the allocation now finds nothing open
    assert allocate.build_problem(store, store.window(w["id"])).dates == []

    # swaps, overrides and reclassification are logged
    d = next(D(k) for k, v in y["shifts"].items() if v["c"] == "AW" and cal.day_type(D(k)) == cal.WEEKDAY)
    Y.swap(store, "2026-27", d, "AW", "XW", "childcare")
    Y.override(store, "2026-27", D("2027-03-31"), None, "nobody after all")
    Y.override(store, "2026-27", D("2027-03-31"), "DL", "DL will do it", kind="LOCUM")
    Y.reclassify(store, "2026-27", d, "JP", "cover arrangement")
    y = store.year("2026-27")
    assert y["shifts"][d.isoformat()] == {"c": "XW", "kind": "JP", "src": "reclassify", "note": "cover arrangement"}
    assert y["shifts"]["2027-03-31"]["c"] == "DL" and y["shifts"]["2027-03-31"]["kind"] == "LOCUM"
    assert [e["event"] for e in y["log"]] == ["allocate", "swap", "override", "override", "reclassify"]

    # export in the workbook's format and read it back
    out = tmp_path / "rota.xlsx"
    xlsx_export.export(store, "2026-27", str(out))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["2026-2027", "Summary"]
    back = xlsx_import.read_year_sheet(wb["2026-2027"], "2026-27")
    assert {k: (v["c"], v["kind"]) for k, v in back["shifts"].items()} == {k: (v["c"], v["kind"]) for k, v in y["shifts"].items()}
    assert [c["initials"] for c in back["consultants"]] == list(PEOPLE)
    assert back["consultants"][0]["weekday"] == 10 and back["consultants"][0]["complete_weekday"] is False  # AW gave one away in the swap
    assert back["consultants"][2]["weekday"] == 12 and back["consultants"][2]["complete_weekday"] is True   # AJ finished
    ws = wb["2026-2027"]
    assert ws["A1"].value == "April" and ws.cell(2, 1).value == 1
    red = [c for row in ws.iter_rows(min_col=2, max_col=8) for c in row if c.fill.fgColor.rgb == "FFFF0000"]
    before_jan = [d for d in cal.daterange(D("2026-04-01"), D("2026-12-31"))]
    expected_red = sum(1 for d in before_jan if cal.day_type(d) != cal.BH)  # bank holidays are yellow, not red
    assert all(isinstance(c.value, int) for c in red) and len(red) == expected_red  # everything before January is unfilled here

    # publish: encrypted payload, no submissions inside, dashboard numbers consistent
    enc = tmp_path / "payload.enc"
    payload = dashboard.publish(store, "2026-27", "correct horse", enc)
    env = json.loads(enc.read_text())
    assert set(env) == {"v", "kdf", "iter", "salt", "iv", "ct"}
    plain = json.loads(dashboard.decrypt(env, "correct horse"))
    assert plain["year"] == "2026-27" and len(plain["rota"]) == 90 and "submissions" not in json.dumps(plain)
    assert [c["initials"] for c in plain["consultants"]] == list(PEOPLE)
    aj = next(c for c in plain["consultants"] if c["initials"] == "AJ")
    assert aj["jp"] == {"weekday": 12, "weekend": 2} and aj["complete"] == {"weekday": True, "weekend": True}
    assert plain["changes"][0]["event"] == "swap" and plain["changes"][0]["reason"] == "childcare"
    m = plain["summary"]["current"]
    assert m["jp_weekday_quota"] == 37 and m["weekday_fill_rate"] == 1.0
    assert m["total_wli_bill"] == costing.wli(store, "2026-27")["committed"]
    import pytest
    with pytest.raises(Exception):
        dashboard.decrypt(env, "wrong passphrase")


def test_sole_offer_front_loads_job_plan_beyond_the_pro_rata_allowance(store):
    add_people(store, {"AW": ("Alan Williams", 10, 2), "XW": ("Ximena Watson", 10, 2)})
    w = W.create(store, "2026-27-Q1", D("2026-04-01"), D("2026-06-30"))
    assert W.windows_left(w) == 4 and W.default_allowance(store, w, R.get(store, "AW")) == {"weekday": 3, "weekend": 1}
    start = D(w["start"])
    aw_days = [D("2026-04-07"), D("2026-04-14"), D("2026-04-21"), D("2026-04-28"), D("2026-05-05"), D("2026-05-12")]
    aw = codec.encode("AW", start, [1 if start + dt.timedelta(days=i) in aw_days else 0 for i in range(W.length(w))], token=w["tokens"]["AW"])
    xw = codec.encode("XW", start, [1 if (start + dt.timedelta(days=i)).weekday() == 2 else 0 for i in range(W.length(w))], token=w["tokens"]["XW"])
    importer.import_text(store, w["id"], aw + "\n" + xw)
    p = allocate.build_problem(store, store.window(w["id"]))
    assert p.allowance["AW"] == {"weekday": 6, "weekend": 1}   # raised from 3: nobody else can do those six Tuesdays
    assert p.allowance["XW"] == {"weekday": 10, "weekend": 1}  # 13 Wednesdays are XW's alone: raised, but never beyond the 10 still owed
    draft = allocate.run(store, w["id"], time_limit=30)
    assert draft["consultants"]["AW"]["jp_window"]["weekday"] == 6 and draft["consultants"]["AW"]["locum_window"] == 0


def test_manual_allowance_override_is_respected(store):
    add_people(store, {"AW": ("Alan Williams", 10, 2)})
    w = W.create(store, "2026-27-Q1", D("2026-04-01"), D("2026-06-30"))
    W.set_allowance(store, w["id"], "AW", weekday=8)
    assert W.allowance(store, store.window(w["id"]), R.get(store, "AW")) == {"weekday": 8, "weekend": 1}
    W.set_allowance(store, w["id"], "AW", weekday=99)
    assert W.allowance(store, store.window(w["id"]), R.get(store, "AW"))["weekday"] == 10  # never more than is owed
