"""Write a job plan year in the workbook's own layout and colours, plus a Summary sheet."""
from __future__ import annotations

import datetime as dt
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter

from . import calendar as cal
from . import roster as R
from . import years as Y
from .dashboard import year_metrics
from .store import Store

GREEN = PatternFill(patternType="solid", fgColor=Color(theme=9, tint=0.5999938962981048))   # locum / WLI
YELLOW = PatternFill(patternType="solid", fgColor="FFFFFF00")                               # bank holiday
RED = PatternFill(patternType="solid", fgColor="FFFF0000")                                  # unfilled
WEEK = PatternFill(patternType="solid", fgColor="FF156082")                                 # week number column
DONE = PatternFill(patternType="solid", fgColor="FF00B050")                                 # job plan complete
MONTH_FONT = Font(bold=True, color="FF0E2841", size=12)
WEEK_FONT = Font(bold=True, color="FFFFFFFF")
BOLD = Font(bold=True)
CENTRE = Alignment(horizontal="center")
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def _week_notes(year: Dict[str, Any]) -> Dict[str, List[str]]:
    notes: Dict[str, List[str]] = {}
    for k, v in (year.get("notes") or {}).items():
        notes.setdefault(k, []).append(v)
    for e in year.get("log", []):
        if e.get("event") in ("swap", "override", "reclassify") and e.get("date"):
            d = dt.date.fromisoformat(e["date"])
            mon = (d - dt.timedelta(days=d.weekday())).isoformat()
            who = e.get("c") or ("%s->%s" % (e.get("from_"), e.get("to")) if e.get("event") == "swap" else (e.get("after") or {}).get("c") if isinstance(e.get("after"), dict) else "")
            notes.setdefault(mon, []).append("%s %s %s: %s" % (d.strftime("%d/%m"), e["event"], who or "", e.get("reason", "")))
    return notes


def write_year_sheet(wb, store: Store, label: str) -> None:
    year = store.year(label)
    ys, ye = cal.year_bounds(label)
    ws = wb.create_sheet(cal.year_label_long(label))
    shifts = year.get("shifts", {})
    unfilled = set(year.get("unfilled", []))
    decided_through = max([dt.date.fromisoformat(k) for k in list(shifts) + list(unfilled)] or [ys - dt.timedelta(days=1)])
    notes = _week_notes(year)
    row = 1
    # month blocks; a week straddling two months is split into two rows, as in the workbook
    d = ys
    week_no = 0
    current_week_monday: Optional[dt.date] = None
    while d <= ye:
        month_first = d
        ws.cell(row, 1, MONTHS[d.month - 1]).font = MONTH_FONT
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        row += 1
        while d <= ye and d.month == month_first.month:
            monday = d - dt.timedelta(days=d.weekday())
            new_week = monday != current_week_monday
            if new_week:
                week_no += 1
                current_week_monday = monday
            a = ws.cell(row, 1, week_no if new_week else "\xa0")
            a.fill = WEEK
            a.font = WEEK_FONT
            a.alignment = CENTRE
            # fill this row with the days of this week that fall in this month
            for col in range(2, 9):
                dd = monday + dt.timedelta(days=col - 2)
                if dd.month != month_first.month or dd < ys or dd > ye:
                    continue
                cell = ws.cell(row, col)
                cell.alignment = CENTRE
                k = dd.isoformat()
                s = shifts.get(k)
                t = cal.day_type(dd)
                if s:
                    cell.value = s["c"]
                    if s["kind"] == Y.LOCUM:
                        cell.fill = GREEN
                    elif s["kind"] == Y.BH or t == cal.BH:
                        cell.fill = YELLOW
                else:
                    cell.value = dd.day
                    if t == cal.BH:
                        cell.fill = YELLOW
                    elif k in unfilled or dd <= decided_through:
                        cell.fill = RED
                if dd.weekday() >= 5 or (s and s["kind"] in (Y.LOCUM, Y.BH)):
                    cell.font = BOLD
            wn = notes.get(monday.isoformat())
            if wn and new_week:
                ws.cell(row, 9, "; ".join(wn))
            # advance d to the first day of the next row
            last_in_row = min(monday + dt.timedelta(days=6), ye)
            if last_in_row.month != month_first.month:
                d = dt.date(month_first.year + (month_first.month // 12), month_first.month % 12 + 1, 1)
            else:
                d = last_in_row + dt.timedelta(days=1)
            row += 1
    # consultant table from column K
    c0 = 11
    heads1 = ["", "", "April-June", "Q1", "July-September", "Q2", "Oct-Dec", "Q3", "Jan-March", "Q4", "Total", "Total", "Job planned weekdays", "Weekends", ""]
    heads2 = ["Consultant", "Initial", "Weekdays", "Weekend days", "Weekdays", "Weekend Days", "Weekdays", "Weekend Days", "Weekdays", "Weekend Days", "Weekdays", "Weekend Days", "", "", "Extra notes"]
    for i, h in enumerate(heads1):
        ws.cell(1, c0 + i, h or None).font = BOLD
    for i, h in enumerate(heads2):
        ws.cell(2, c0 + i, h or None).font = BOLD
    counts = Y.counts(year)
    per_q: Dict[str, Dict[int, Dict[str, int]]] = {}
    for k, s in shifts.items():
        dd = dt.date.fromisoformat(k)
        q = per_q.setdefault(s["c"], {1: {"weekday": 0, "weekend": 0}, 2: {"weekday": 0, "weekend": 0}, 3: {"weekday": 0, "weekend": 0}, 4: {"weekday": 0, "weekend": 0}})
        t = "weekend" if dd.weekday() >= 5 else "weekday"
        q[cal.quarter_of(dd)][t] += 1
    r = 3
    tot = {"wd": 0, "we": 0, "qwd": 0, "qwe": 0}
    for c in R.consultants(store):
        ini = c["initials"]
        qw, qe = R.quota(c, label)
        pq = per_q.get(ini, {n: {"weekday": 0, "weekend": 0} for n in range(1, 5)})
        ws.cell(r, c0, c["name"])
        ws.cell(r, c0 + 1, ini)
        for qn in range(1, 5):
            ws.cell(r, c0 + 2 + (qn - 1) * 2, pq[qn]["weekday"])
            ws.cell(r, c0 + 3 + (qn - 1) * 2, pq[qn]["weekend"])
        twd = sum(pq[qn]["weekday"] for qn in range(1, 5))
        twe = sum(pq[qn]["weekend"] for qn in range(1, 5))
        ws.cell(r, c0 + 10, twd)
        ws.cell(r, c0 + 11, twe)
        k = counts.get(ini, {"jp_weekday": 0, "jp_weekend": 0})
        cq = ws.cell(r, c0 + 12, qw)
        ce = ws.cell(r, c0 + 13, qe)
        if k["jp_weekday"] >= qw:
            cq.fill = DONE
        if k["jp_weekend"] >= qe:
            ce.fill = DONE
        ws.cell(r, c0 + 14, c.get("notes") or None)
        tot["wd"] += twd
        tot["we"] += twe
        tot["qwd"] += qw
        tot["qwe"] += qe
        r += 1
    r += 1
    ws.cell(r, c0 + 10, tot["wd"]).font = BOLD
    ws.cell(r, c0 + 11, tot["we"]).font = BOLD
    ws.cell(r, c0 + 12, tot["qwd"]).font = BOLD
    ws.cell(r, c0 + 13, tot["qwe"]).font = BOLD
    r += 1
    ws.cell(r, c0, "Vacant")
    ws.cell(r, c0 + 1).fill = RED
    ws.cell(r, c0 + 12, "Job planned completed")
    ws.cell(r, c0 + 13).fill = DONE
    r += 1
    ws.cell(r, c0, "Bank Holiday")
    ws.cell(r, c0 + 1, "Bank").fill = YELLOW
    r += 1
    ws.cell(r, c0, "Extra Shift")
    ws.cell(r, c0 + 1, "Bank").fill = GREEN
    for col, w in zip(range(1, 10), [6, 7, 7, 7, 7, 7, 7, 7, 40]):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.column_dimensions[get_column_letter(c0)].width = 25
    for i in range(1, 14):
        ws.column_dimensions[get_column_letter(c0 + i)].width = 14
    ws.column_dimensions[get_column_letter(c0 + 14)].width = 50


def write_summary_sheet(wb, store: Store, label: str) -> None:
    ws = wb.create_sheet("Summary")
    hist = store.history().get("years", {})
    cur = year_metrics(store, label)
    cols = sorted(set(hist) | {label})
    rows = [
        ("JP weekdays shifts", "jp_weekday_quota", None), ("JP weekend day shift", "jp_weekend_quota", None), ("JP bank holiday", "jp_bank_holiday", None),
        ("Bank available", "budget", "£#,##0"), (None, None, None),
        ("Weekday fill rate", "weekday_fill_rate", "0%"), ("Weekend fill rate", "weekend_fill_rate", "0%"), ("Bank holiday fill rate", "bank_holiday_fill_rate", "0%"), (None, None, None),
        ("Weekday fill rate above JP", "weekday_fill_rate_above_jp", "0%"), ("Weekend fill rate above JP", "weekend_fill_rate_above_jp", "0%"), (None, None, None),
        ("Total WLI bill", "total_wli_bill", "£#,##0"),
    ]
    for i, lab in enumerate(cols):
        ws.cell(1, 2 + i, cal.year_label_long(lab)).font = BOLD
    for ri, (title, key, fmt) in enumerate(rows, start=2):
        if not title:
            continue
        ws.cell(ri, 1, title)
        for i, lab in enumerate(cols):
            m = cur if lab == label else hist.get(lab, {})
            v = m.get(key)
            if v is None:
                continue
            cell = ws.cell(ri, 2 + i, v)
            if fmt:
                cell.number_format = fmt
    ws.column_dimensions["A"].width = 30
    for i in range(len(cols)):
        ws.column_dimensions[get_column_letter(2 + i)].width = 14


def export(store: Store, label: str, out_path: str) -> str:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    write_year_sheet(wb, store, label)
    write_summary_sheet(wb, store, label)
    wb.save(out_path)
    return out_path
