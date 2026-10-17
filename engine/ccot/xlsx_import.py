"""Read the existing rota workbook: roster and quotas, this year's allocations, and the Summary tab.

Conventions found in the workbook (2025-2026 and 2026-2027 sheets):
  column A  week number (new week) or a non-breaking space (the same week continued into the next month block)
  columns B-H  Monday to Sunday; initials = filled, bare day number = not filled
  fills  green (theme accent6, tint 0.6) = locum / WLI, yellow = bank holiday, red = unfilled, none = job plan
  column I  free-text notes for that week
  from column K  the consultant table: name, initials, quarterly counts, totals, quotas, notes;
                 a green quota cell means the job plan is complete
"""
from __future__ import annotations

import datetime as dt
import re
from typing import Any, Dict, List, Optional, Tuple

import openpyxl

from . import calendar as cal
from . import years as Y

MONTHS = {m.lower() for m in ("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December")}


def _fill(cell) -> str:
    f = cell.fill
    if f is None or f.fill_type is None:
        return ""
    fg = f.fgColor
    if fg.type == "rgb":
        return "rgb:%s" % fg.rgb
    if fg.type == "theme":
        return "theme:%d:%.2f" % (fg.theme, fg.tint or 0.0)
    return "other"


def kind_from_fill(fill: str) -> Optional[str]:
    """None for no fill (job plan), or LOCUM, BH, UNFILLED, or 'unknown'."""
    if fill in ("", "rgb:FFFFFFFF", "theme:0:0.00"):
        return None
    if fill.startswith("theme:9:0.6") or fill == "rgb:FFC6E0B4":
        return Y.LOCUM
    if fill == "rgb:FFFFFF00":
        return Y.BH
    if fill == "rgb:FFFF0000":
        return "UNFILLED"
    return "unknown"


def _is_month_header(v: Any) -> bool:
    return isinstance(v, str) and any(part.strip().lower() in MONTHS for part in v.split("/"))


def read_year_sheet(ws, label: str) -> Dict[str, Any]:
    ys, _ = cal.year_bounds(label)
    week0_monday = ys - dt.timedelta(days=ys.weekday())
    out: Dict[str, Any] = {"shifts": {}, "unfilled": [], "notes": {}, "anomalies": [], "consultants": [], "side_totals": {}}
    anomalies = out["anomalies"]
    week = -1
    prev_full = False
    seen: Dict[dt.date, Tuple[int, int]] = {}
    bh_open: List[dt.date] = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        row_vals = [ws.cell(r, c).value for c in range(2, 9)]
        has_grid = any(v not in (None, "") for v in row_vals)
        if _is_month_header(a):
            continue
        if isinstance(a, (int, float)) or (isinstance(a, str) and re.fullmatch(r"\d+(st|nd|rd|th)?", a.strip())):
            week += 1
        elif has_grid and (a is None or (isinstance(a, str) and a.strip() == "")):
            if prev_full:
                week += 1
        else:
            continue
        prev_full = all(v not in (None, "") for v in row_vals)
        if week < 0:
            continue
        for c in range(2, 9):
            cell = ws.cell(r, c)
            v = cell.value
            if v is None or (isinstance(v, str) and v.strip() in ("", "\xa0")):
                continue
            d = week0_monday + dt.timedelta(days=7 * week + (c - 2))
            if d in seen:
                anomalies.append("%s appears twice (rows %d and %d)" % (d, seen[d][0], r))
                continue
            seen[d] = (r, c)
            fill = _fill(cell)
            kind = kind_from_fill(fill)
            if isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit()):
                num = int(float(v))
                if num != d.day:
                    anomalies.append("cell %s%d shows day %d but the date is %s" % (cell.column_letter, r, num, d))
                if kind == "UNFILLED":
                    out["unfilled"].append(d.isoformat())
                elif kind == Y.BH and ys <= d:
                    bh_open.append(d)   # unfilled only if inside the period already allocated
                continue
            ini = str(v).strip().upper()
            if "/" in ini:
                anomalies.append("cell %s%d has a split day %r; recorded as %s" % (cell.column_letter, r, v, ini.split("/")[0]))
                ini = ini.split("/")[0].strip()
            if not re.fullmatch(r"[A-Z]{1,3}", ini):
                anomalies.append("cell %s%d has unexpected text %r; ignored" % (cell.column_letter, r, v))
                continue
            if kind == "unknown":
                anomalies.append("cell %s%d (%s, %s) has an unrecognised fill %s; treated as job plan" % (cell.column_letter, r, d, ini, fill))
                kind = None
            if kind == "UNFILLED":
                anomalies.append("cell %s%d (%s) is red but holds %s; treated as job plan" % (cell.column_letter, r, d, ini))
                kind = None
            if cal.day_type(d) == cal.BH and kind != Y.BH:
                anomalies.append("%s is a bank holiday but the cell is not yellow; recorded as a bank holiday shift" % d)
                kind = Y.BH
            if kind == Y.BH and cal.day_type(d) != cal.BH:
                anomalies.append("cell %s%d (%s, %s) is yellow but %s is not a bank holiday; treated as locum" % (cell.column_letter, r, d, ini, d))
                kind = Y.LOCUM
            out["shifts"][d.isoformat()] = {"c": ini, "kind": kind or Y.JP, "src": "import", "note": None}
        note = ws.cell(r, 9).value
        if isinstance(note, str) and note.strip():
            monday = week0_monday + dt.timedelta(days=7 * week)
            out["notes"][monday.isoformat()] = (out["notes"].get(monday.isoformat(), "") + " " + note.strip()).strip()
    decided_through = max([dt.date.fromisoformat(k) for k in list(out["shifts"]) + out["unfilled"]] or [ys - dt.timedelta(days=1)])
    out["unfilled"].extend(d.isoformat() for d in bh_open if d <= decided_through)
    out["unfilled"] = sorted(set(out["unfilled"]))
    out["decided_through"] = decided_through.isoformat()
    # consultant table
    hdr = None
    for r in range(1, min(ws.max_row, 10) + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().lower() == "consultant":
                hdr = (r, c)
                break
        if hdr:
            break
    if hdr:
        r0, c0 = hdr
        cols = {}
        for c in range(c0, ws.max_column + 1):
            for rr in (r0 - 1, r0):
                v = ws.cell(rr, c).value
                if isinstance(v, str):
                    t = v.strip().lower()
                    if t.startswith("job planned weekday"):
                        cols["qwd"] = c
                    elif t == "weekends" and "qwe" not in cols and "qwd" in cols:
                        cols["qwe"] = c
                    elif t.startswith("extra notes"):
                        cols["notes"] = c
        r = r0 + 1
        while r <= ws.max_row:
            name = ws.cell(r, c0).value
            ini = ws.cell(r, c0 + 1).value
            if not (isinstance(name, str) and name.strip()) or not (isinstance(ini, str) and ini.strip()):
                break
            qwd = ws.cell(r, cols["qwd"]).value if "qwd" in cols else None
            qwe = ws.cell(r, cols["qwe"]).value if "qwe" in cols else None
            done_wd = "qwd" in cols and _fill(ws.cell(r, cols["qwd"])) == "rgb:FF00B050"
            done_we = "qwe" in cols and _fill(ws.cell(r, cols["qwe"])) == "rgb:FF00B050"
            note = ws.cell(r, cols["notes"]).value if "notes" in cols else None
            out["consultants"].append({
                "name": name.strip(), "initials": ini.strip().upper(),
                "weekday": int(qwd) if isinstance(qwd, (int, float)) else None,
                "weekend": int(qwe) if isinstance(qwe, (int, float)) else None,
                "complete_weekday": done_wd, "complete_weekend": done_we,
                "notes": note.strip() if isinstance(note, str) else "",
            })
            r += 1
    # cross-checks
    known = {c["initials"] for c in out["consultants"]}
    counts: Dict[str, Dict[str, int]] = {}
    for k, s in out["shifts"].items():
        d = dt.date.fromisoformat(k)
        cc = counts.setdefault(s["c"], {"jp_weekday": 0, "jp_weekend": 0, "locum_weekday": 0, "locum_weekend": 0, "bh": 0})
        if s["kind"] == Y.BH:
            cc["bh"] += 1
        else:
            cc[("jp_" if s["kind"] == Y.JP else "locum_") + cal.day_type(d)] += 1
    for ini in sorted(counts):
        if ini not in known:
            anomalies.append("%s appears in the grid but not in the consultant table" % ini)
    for c in out["consultants"]:
        cc = counts.get(c["initials"], {"jp_weekday": 0, "jp_weekend": 0, "locum_weekday": 0, "locum_weekend": 0, "bh": 0})
        for t in ("weekday", "weekend"):
            q = c[t]
            if q is None:
                continue
            if cc["jp_" + t] > q:
                anomalies.append("%s has %d white (job plan) %s shifts against a quota of %d" % (c["initials"], cc["jp_" + t], t, q))
            if cc["locum_" + t] and cc["jp_" + t] < q:
                anomalies.append("%s has %d green (locum) %s shifts while only %d of %d job plan %s shifts are white" % (c["initials"], cc["locum_" + t], t, cc["jp_" + t], q, t))
    out["counts"] = counts
    return out


def read_summary(ws) -> Dict[str, Dict[str, Any]]:
    """Summary tab -> {year label: {metric: value}}."""
    keys = {
        "jp weekdays shifts": "jp_weekday_quota", "jp weekend day shift": "jp_weekend_quota", "jp bank holiday": "jp_bank_holiday",
        "bank available": "budget", "weekday fill rate": "weekday_fill_rate", "weekend fill rate": "weekend_fill_rate",
        "bank holiday fill rate": "bank_holiday_fill_rate", "weekday fill rate above jp": "weekday_fill_rate_above_jp",
        "weekend fill rate above jp": "weekend_fill_rate_above_jp", "total wli bill": "total_wli_bill",
    }
    years: Dict[int, str] = {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str) and re.fullmatch(r"\d{4}-\d{4}", v.strip()):
            years[c] = cal.label_from_long(v.strip())
    out: Dict[str, Dict[str, Any]] = {lab: {} for lab in years.values()}
    for r in range(2, ws.max_row + 1):
        lab = ws.cell(r, 1).value
        if not isinstance(lab, str):
            continue
        key = keys.get(lab.strip().lower())
        if not key:
            continue
        for c, year in years.items():
            v = ws.cell(r, c).value
            if isinstance(v, str):
                s = v.strip().replace("%%", "%")
                if s.endswith("%"):
                    try:
                        v = float(s[:-1]) / 100.0
                    except ValueError:
                        v = None
                else:
                    v = None
            out[year][key] = v
    return {k: v for k, v in out.items() if any(x is not None for x in v.values())}


def read_workbook(path: str, label: str) -> Dict[str, Any]:
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = cal.year_label_long(label)
    if sheet not in wb.sheetnames:
        raise ValueError("no sheet %r in %s (sheets: %s)" % (sheet, path, ", ".join(wb.sheetnames)))
    data = read_year_sheet(wb[sheet], label)
    data["history"] = read_summary(wb["Summary"]) if "Summary" in wb.sheetnames else {}
    return data


def import_workbook(store, path: str, label: str, replace_shifts: bool = False) -> Dict[str, Any]:
    """Seed the data folder from the workbook. Returns the anomaly report."""
    from . import roster as R

    data = read_workbook(path, label)
    r = store.roster()
    existing = {c["initials"]: c for c in r["consultants"]}
    for c in data["consultants"]:
        has_shifts = c["initials"] in data["counts"]
        active = has_shifts or bool(c["weekday"]) or bool(c["weekend"])
        if c["initials"] in existing:
            e = existing[c["initials"]]
            e["name"] = c["name"]
            e.setdefault("quotas", {})[label] = {"weekday": c["weekday"] or 0, "weekend": c["weekend"] or 0}
        else:
            r["consultants"].append({"initials": c["initials"], "name": c["name"], "active": active,
                                     "quotas": {label: {"weekday": c["weekday"] or 0, "weekend": c["weekend"] or 0}},
                                     "notes": c["notes"]})
    store.save_roster(r)
    y = store.year(label)
    if replace_shifts:
        y["shifts"] = {}
        y["unfilled"] = []
    for k, s in data["shifts"].items():
        if k not in y["shifts"]:
            y["shifts"][k] = s
    y["unfilled"] = sorted(set(y.get("unfilled", [])) | set(data["unfilled"]) - set(y["shifts"]))
    y.setdefault("notes", {}).update(data["notes"])
    Y.log(y, "import", source=str(path), shifts=len(data["shifts"]), unfilled=len(data["unfilled"]), anomalies=len(data["anomalies"]))
    store.save_year(y)
    h = store.history()
    for lab, metrics in data["history"].items():
        h.setdefault("years", {}).setdefault(lab, {}).update({k: v for k, v in metrics.items() if v is not None})
    store.save_history(h)
    s = store.settings()
    for lab, metrics in data["history"].items():
        if metrics.get("budget"):
            s.setdefault("budgets", {}).setdefault(lab, metrics["budget"])
    store.save_settings(s)
    return {"consultants": len(data["consultants"]), "shifts": len(data["shifts"]), "unfilled": data["unfilled"],
            "anomalies": data["anomalies"], "history_years": sorted(data["history"]), "counts": data["counts"]}
